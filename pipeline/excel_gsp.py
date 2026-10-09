"""
Builds the two "ready for humans" Excel exports derived from the raw MySQL schedule data:

* GSP.xlsx - the full schedule for every team, one row per game, colour-coded per team,
  with a "KG" (Kampfgericht / scorer's-table duty) column filled in.
* KG.xlsx  - the same data trimmed down to what the scorer's-table duty roster needs,
  sorted by date and colour-coded by home/away.

This mirrors the original script's output exactly - just split into small functions with
docstrings instead of one long procedural block.

Run directly: ``python send2excel.py``
"""
from __future__ import annotations

import logging

import openpyxl
import openpyxl.styles
import pandas as pd
from openpyxl.styles import Alignment, Border, Side, numbers
from openpyxl.worksheet.worksheet import Worksheet

from pipeline import db
from pipeline.kampfgericht import KAMPFGERICHT_ASSIGNMENTS
from pipeline.logging_utils import configure_logging
from pipeline.settings import SETTINGS
from pipeline.teams import HOME_TEAM_ALIASES, TEAM_COLORS

logger = logging.getLogger(__name__)

# All team tables, including the ones not yet tracked by the scraper (their tables just
# stay empty, which is fine - SELECT * FROM an empty table simply contributes no rows).
TEAM_TABLES = [
    "h1", "d1", "h2", "u18", "u18w", "u16", "u16w", "u141", "u142", "u14w", "u121", "u122", "u10",
]

COLUMN_WIDTH = 30
BORDER = Border(
    left=Side(border_style="thin", color="000000"),
    right=Side(border_style="thin", color="000000"),
    top=Side(border_style="thin", color="000000"),
    bottom=Side(border_style="thin", color="000000"),
)
BOLD_FONT = openpyxl.styles.Font(bold=True, size=14)
HEADING_FILL = openpyxl.styles.PatternFill(patternType="solid", fgColor="CCCCFF")
HOME_GAME_FILL = openpyxl.styles.PatternFill(patternType="solid", fgColor="BDD7EE")
AWAY_GAME_FILL = openpyxl.styles.PatternFill(patternType="solid", fgColor="BFBFBF")


def _assign_kampfgericht(df: pd.DataFrame) -> pd.DataFrame:
    """Fills the 'KG' column using the hand-curated game-id -> team-tag assignments."""
    for game_id, team_tag in KAMPFGERICHT_ASSIGNMENTS.items():
        df.loc[df["id"] == game_id, "KG"] = team_tag
    return df


def build_gsp_dataframe() -> pd.DataFrame:
    """Loads every team's games and attaches the Kampfgericht duty column."""
    df = db.fetch_games_df(TEAM_TABLES)
    df.insert(12, "KG", "")
    return _assign_kampfgericht(df)


def _apply_common_table_styles(sheet: Worksheet, column_letters: list[str]) -> None:
    """Shared styling for both GSP and KG sheets: borders, header row, column widths, filters."""
    for column in column_letters:
        sheet.column_dimensions[column].width = COLUMN_WIDTH

    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal="right")
            cell.border = BORDER

    for cell in sheet[1]:
        cell.font = BOLD_FONT
        cell.fill = HEADING_FILL
        cell.alignment = Alignment(horizontal="center")

    sheet.auto_filter.ref = sheet.dimensions


def style_gsp_workbook(path) -> None:
    """Colour-codes every row by team, formats the kickoff-time column, adds filters."""
    workbook = openpyxl.load_workbook(path)
    sheet = workbook.active

    for cell in sheet["F"]:
        cell.number_format = numbers.FORMAT_DATE_TIME3

    column_letters = ["A", "B", "C", "D", "E", "F", "G", "H", "I"]
    _apply_common_table_styles(sheet, column_letters)

    team_column_index = 2  # 0-based index of the "team" column within each row
    for row in range(2, sheet.max_row + 1):
        team_tag = sheet.cell(row=row, column=team_column_index + 1).value
        color = TEAM_COLORS.get(team_tag, "FFFFFF")
        fill = openpyxl.styles.PatternFill(patternType="solid", fgColor=color)
        for cell in sheet[row]:
            cell.fill = fill

    workbook.save(path)


def build_kg_dataframe(gsp_df: pd.DataFrame) -> pd.DataFrame:
    """Derives the Kampfgericht roster view: fewer columns, sorted by date."""
    kg_df = gsp_df.copy()
    kg_df = kg_df.drop(columns=["gast_kurz", "heim_kurz", "liga", "spnr", "spt"])
    kg_df = _assign_kampfgericht(kg_df)
    return kg_df.sort_values(by="datum", ascending=True)


def style_kg_workbook(path) -> None:
    """Formats the kickoff-time column and colours rows by home/away game."""
    workbook = openpyxl.load_workbook(path)
    sheet = workbook.active

    for cell in sheet["C"]:
        cell.number_format = numbers.FORMAT_DATE_TIME3

    column_letters = ["A", "B", "C", "D", "E", "F", "G", "H", "I"]
    _apply_common_table_styles(sheet, column_letters)

    home_team_column_index = 3  # 0-based index of the "heim" column within each row
    for row in range(2, sheet.max_row + 1):
        home_team = sheet.cell(row=row, column=home_team_column_index + 1).value
        fill = HOME_GAME_FILL if home_team in HOME_TEAM_ALIASES else AWAY_GAME_FILL
        for cell in sheet[row]:
            cell.fill = fill

    workbook.save(path)


def run() -> None:
    """Builds, styles and saves both GSP.xlsx and KG.xlsx to the Google Drive folder."""
    configure_logging()
    SETTINGS.output_dir.mkdir(parents=True, exist_ok=True)

    gsp_df = build_gsp_dataframe()
    gsp_path = SETTINGS.output_dir / "GSP.xlsx"
    gsp_df.to_excel(gsp_path, index=False)
    style_gsp_workbook(gsp_path)
    logger.info("GSP export written to %s (%d games)", gsp_path, len(gsp_df))

    kg_df = build_kg_dataframe(gsp_df)
    kg_path = SETTINGS.output_dir / "KG.xlsx"
    kg_df.to_excel(kg_path, index=False)
    style_kg_workbook(kg_path)
    logger.info("Kampfgericht export written to %s (%d games)", kg_path, len(kg_df))


if __name__ == "__main__":
    run()

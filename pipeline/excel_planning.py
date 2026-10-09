"""
Builds the master "planning table" Excel sheet: one row per day in the planning window,
one column per team, showing whether/when that team plays and against whom. Weekends,
school holidays and blocked dates (tournaments etc.) are highlighted to make manual hall
slot planning easier.

This mirrors the original script's output exactly - same layout, same colours - just
split into small functions with docstrings instead of one long procedural block.

Run directly: ``python createPlanningTable.py``
"""
from __future__ import annotations

import logging
from datetime import timedelta

import openpyxl
import openpyxl.styles
import pandas as pd
from openpyxl.styles import Alignment, Border, Side
from openpyxl.worksheet.worksheet import Worksheet

from pipeline import db
from pipeline.logging_utils import configure_logging
from pipeline.settings import ADDITIONAL_EVENTS, HOLIDAYS, PLANNING_END, PLANNING_START, SETTINGS
from pipeline.teams import HOME_TEAM_ALIASES, OPPONENT_ABBREVIATIONS, TEAM_TABLES

logger = logging.getLogger(__name__)

LOCAL_WORKING_FILE = "planningTable.xlsx"
COLUMN_WIDTH = 20

BORDER = Border(
    left=Side(border_style="thin", color="000000"),
    right=Side(border_style="thin", color="000000"),
    top=Side(border_style="thin", color="000000"),
    bottom=Side(border_style="thin", color="000000"),
)
BOLD_FONT = openpyxl.styles.Font(bold=True, size=15)
WEEKEND_FILL = openpyxl.styles.PatternFill(patternType="solid", fgColor="969696")
HOLIDAY_FILL = openpyxl.styles.PatternFill(patternType="solid", fgColor="CCFFCC")
BLOCKED_FILL = openpyxl.styles.PatternFill(patternType="solid", fgColor="993300")
HEADING_FILL = openpyxl.styles.PatternFill(patternType="solid", fgColor="CCCCFF")


def _all_planning_dates() -> list[str]:
    """Returns every date in the planning window as 'YYYY-MM-DD' strings."""
    dates = []
    current = PLANNING_START
    while current <= PLANNING_END:
        dates.append(current.strftime("%Y-%m-%d"))
        current += timedelta(days=1)
    return dates


def build_pivot_table() -> pd.DataFrame:
    """Builds the day x team pivot: each cell is '<kickoff time> vs/@ <opponent>'."""
    games = db.fetch_games_df(TEAM_TABLES)
    pivot = pd.DataFrame(index=_all_planning_dates(), columns=TEAM_TABLES)

    for _, game in games.iterrows():
        kickoff = str(game["uhrzeit"])[:-3]  # drop trailing ":00" seconds
        if game["heim"] in HOME_TEAM_ALIASES:
            opponent = OPPONENT_ABBREVIATIONS.get(game["gast"], game["gast"])
            label = f"{kickoff} vs {opponent}"
        else:
            opponent = OPPONENT_ABBREVIATIONS.get(game["heim"], game["heim"])
            label = f"{kickoff} @ {opponent}"
        pivot.loc[str(game["datum"]), game["team"]] = label

    return pivot


def _apply_base_styles(sheet: Worksheet) -> None:
    """Right-aligns and borders every cell, bolds/colours the header row and column."""
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal="right")
            cell.border = BORDER

    for cell in sheet["A"]:
        cell.font = BOLD_FONT
        cell.fill = HEADING_FILL

    for cell in sheet[1]:
        cell.font = BOLD_FONT
        cell.fill = HEADING_FILL
        cell.alignment = Alignment(horizontal="centerContinuous")

    for column in "ABCDEFGHIJK":
        sheet.column_dimensions[column].width = COLUMN_WIDTH


def _highlight_weekends(sheet: Worksheet) -> None:
    """Greys out every Saturday/Sunday row pair (7 rows per week, starting row 2)."""
    for row_index in range(2, 250, 7):
        for cell in sheet[row_index]:
            cell.fill = WEEKEND_FILL
        for cell in sheet[row_index + 1]:
            cell.fill = WEEKEND_FILL


def _highlight_holidays(sheet: Worksheet) -> None:
    """Marks school holiday dates in column N and shades the corresponding 2-day block."""
    for holiday_date in HOLIDAYS:
        for row_index, row_values in enumerate(sheet.iter_rows(values_only=True), start=1):
            if holiday_date not in row_values:
                continue
            sheet[f"N{row_index}"] = "Schulferien"
            for block_row in range(row_index, row_index + 2):
                for cell in sheet[block_row]:
                    if cell.column_letter == "N":
                        cell.font = BOLD_FONT
                        break
                    cell.fill = HOLIDAY_FILL
            break


def _highlight_additional_events(sheet: Worksheet) -> None:
    """Marks blocked dates (tournaments etc.) in column N and shades the blocked rows."""
    for event_name, (event_date, blocked) in ADDITIONAL_EVENTS.items():
        for row_index, row_values in enumerate(sheet.iter_rows(values_only=True), start=1):
            if event_date not in row_values:
                continue

            if isinstance(blocked, int):
                # `blocked` is a number of consecutive days the hall is unavailable
                sheet[f"N{row_index}"] = event_name
                for block_row in range(row_index, row_index + blocked):
                    for cell in sheet[block_row]:
                        if cell.column_letter == "N":
                            cell.font = BOLD_FONT
                            break
                        cell.fill = BLOCKED_FILL
            else:
                # `blocked` is a single "HH:MM" time slot on that one day
                sheet[f"N{row_index}"] = f"{event_name} um {blocked}"
                for cell in sheet[row_index]:
                    if cell.column_letter == "N":
                        cell.font = BOLD_FONT
                        break
                    cell.fill = BLOCKED_FILL
            break


def style_workbook(path: str) -> None:
    """Opens the given workbook and applies all highlighting/styling in place."""
    workbook = openpyxl.load_workbook(path)
    sheet = workbook.active

    _apply_base_styles(sheet)
    _highlight_weekends(sheet)
    _highlight_holidays(sheet)
    _highlight_additional_events(sheet)

    workbook.save(path)


def run() -> None:
    """Builds and styles the planning table, saving a local copy and the Drive copy."""
    configure_logging()

    pivot = build_pivot_table()
    pivot.to_excel(LOCAL_WORKING_FILE)
    style_workbook(LOCAL_WORKING_FILE)

    output_path = SETTINGS.output_dir / "planningTable.xlsx"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    openpyxl.load_workbook(LOCAL_WORKING_FILE).save(output_path)

    logger.info("Planning table written to %s", output_path)


if __name__ == "__main__":
    run()
